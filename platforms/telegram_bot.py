"""Telegram front end for the agent.

Open to anyone. The owner (OWNER_TELEGRAM_ID) has no limits; everyone else gets
DAILY_LIMIT requests per day, an onboarding chat, and personalised news.
"""
import asyncio
import base64
import re
import logging
import os
import html
import secrets
import tempfile
import time
import types

import fitz
import openpyxl
from docx import Document
from dotenv import load_dotenv
from pptx import Presentation
from telegram import (InlineKeyboardButton, InlineKeyboardMarkup, InlineQueryResultArticle,
                      InlineQueryResultsButton, InputTextMessageContent, Update)
from telegram.constants import ChatAction
from telegram.ext import (ApplicationBuilder, CallbackQueryHandler, ChatMemberHandler, ChosenInlineResultHandler,
                          CommandHandler, ContextTypes, InlineQueryHandler, MessageHandler, filters)

from core import context as agent_context
from platforms import tg_live
from core import alerts, downloader, llm, models, profile, store, topics, voice
from core.agent import chat
chat_fn = chat
from core.memory import Memory

load_dotenv()
logging.basicConfig(level=logging.WARNING)

memory = Memory()
client = llm.Client("vision")
DAILY_LIMIT = int(os.getenv("DAILY_LIMIT", "20"))
VISION_OWNER_MODEL = "claude-sonnet-5-5"
VISION_PUBLIC_MODEL = "claude-haiku-4-5-20251001"

FIRST_HELLO = ("Hi{name}! I'm an AI assistant: ask me anything, send links, files, photos or voice notes. "
               "Tip: /help shows what I can do. Chats are saved on the bot owner's hard disk, and the owner can read them.")

WELCOME = (
    "Hi{name}! I'm an AI assistant built by Naveen Raj.\n\n"
    "I can answer questions, search the web, read links and YouTube videos, download videos from X, Instagram "
    "and TikTok, read your files and images, understand voice notes, set reminders, and send you a personalised "
    "news brief every morning and evening.\n\n"
    "Just ask me anything to get started. Below you can pick topics you want updates on.\n\n"
    "I learn from our chats to personalise answers and news. Chats are saved on the bot owner's hard disk, and the "
    "owner can read them to improve the bot. See what I know with /profile, erase it all with /forget."
)

PRIVACY = "\nPrivacy: chats are saved on the bot owner's hard disk, and the owner can read them."
HELP = (
    "Ask me anything, like you would ChatGPT, Claude, Gemini or Perplexity: questions, research with sources, "
    "writing, studying, coding, maths, advice. Send links, PDFs, Word/Excel/PowerPoint files, photos or voice notes "
    "and I'll read them. Send any link and I'll tell you what it is; add the word \"download\" to a video link "
    "(X, Instagram, TikTok, YouTube, Reddit) to get the file.\n\n"
    "In any chat, type @{bot} and your question to post an answer right there.\n\n"
    "/new - start a fresh conversation\n"
    "/search <question> - web answer with sources\n"
    "/interests - pick topics you want updates on (or type your own)\n"
    "/updates - how often: live, twice a day, mornings, or off\n"
    "/news - your personalised news now\n"
    "/keywords <words> - track exact keywords across news, Reddit, X, Instagram, YouTube\n"
    "/newsoff, /newson - stop or restart updates\n"
    "/channels add @name - follow a public Telegram channel in your updates\n"
    "/alert <keyword> - instant alerts when something new appears\n"
    "/alerts - your alerts, /unalert <id> to remove\n"
    "/profile - what I know about you\n"
    "/forget - erase what I know about you and our chats\n"
    "/model - see models, set your default\n"
    "/switch - pick a different model\n"
    "/picker on|off - ask which model at the start of each conversation\n"
    "/style <how you like replies> - e.g. short, casual, Hinglish\n"
    "/clear - reset this conversation\n"
    "/deep <question> - strongest model (owner)\n"
    "/quick <question> - fastest model\n"
)
OWNER_HELP = HELP + (
    "\nOwner only:\n"
    "/brief - full strategic briefing now\n"
    "/channels brief - what's new in the channels you've joined\n"
    "/event - tournament viewership (BMSD): peak, average, hours watched; /event help\n"
    "/radar - content and business radar now\n"
    "/users - who is using the bot today\n"
    "/status - health check: sources, spend, scheduler\n"
    "/usage - what the bot cost today and this week\n"
    "/credits <amount> - track your Anthropic balance\n"
    "/dashboard - link to the web dashboard\n"
    "/scout @channel - creator numbers + fit, saved to tracker\n"
    "/script <topic> - video script in your brand voice\n"
    "/shorts <YouTube link> - 5 clip moments with timestamps\n"
    "/captions <topic> - caption pack for every platform\n"
    "/brandvoice add <sample> - teach it your writing style\n"
    "/groups, /leave <id> - groups I'm in\n"
    "/backup - back up the database now\n"
    "/pipeline - talent agency pipeline\n"
    "/activity [user id] - what people asked and searched\n"
    "/thinking off|low|medium|high - how hard models think\n"
    "/think <question> - think hard for one message\n"
    "/block <id>, /unblock <id> - moderate users\n"
)


# ---------- helpers ----------

def action_kb(ctx):
    """Link buttons the agent's tools asked for (Open in Gmail, Add to Calendar, Send on WhatsApp)."""
    if not getattr(ctx, "buttons", None):
        return None
    rows, row = [], []
    for label, url in ctx.buttons[:8]:
        row.append(InlineKeyboardButton(label, url=url))
        if len(row) == 2:
            rows.append(row)
            row = []
    if row:
        rows.append(row)
    return InlineKeyboardMarkup(rows)


def links_kb(uid: str):
    from core import newslinks
    key = newslinks.last_key(uid)
    return InlineKeyboardMarkup([[InlineKeyboardButton("Links", callback_data=f"nl:{key}")]]) if key else None


async def on_news_links(update: Update, context: ContextTypes.DEFAULT_TYPE):
    from core import newslinks
    q = update.callback_query
    uid = store.uid_for("telegram", str(update.effective_user.id))
    await q.answer()
    text = newslinks.links_text(uid, q.data.split(":", 1)[1])
    for i in range(0, len(text), 4000):
        await q.message.reply_text(text[i:i + 4000], disable_web_page_preview=True)


def _bot_name(text: str, context) -> str:
    return text.replace("{bot}", getattr(context.bot, "username", None) or "this bot's name")


def ids(update: Update) -> tuple:
    u = update.effective_user
    return str(u.id), str(update.effective_chat.id), (u.first_name or ""), (u.username or "")


def is_owner(update: Update) -> bool:
    return store.is_owner("telegram", str(update.effective_user.id))


async def send_long(update: Update, text: str, rich: bool = False, markup=None, preview: bool = True):
    """Send a reply in Telegram-sized pieces. rich=True renders Markdown (bold, lists, code, links).
    markup (buttons) goes under the last piece."""
    text = text or "(no response)"
    if not rich:
        parts = [text[i:i + 4096] for i in range(0, len(text), 4096)]
        for n, part in enumerate(parts):
            extra = {"reply_markup": markup} if markup and n == len(parts) - 1 else {}
            if not preview:
                extra["disable_web_page_preview"] = True
            await update.message.reply_text(part, **extra)
        return
    from core import tg_format
    chunks = tg_format.split_for_telegram(text)
    for n, chunk in enumerate(chunks):
        extra = {"reply_markup": markup} if markup and n == len(chunks) - 1 else {}
        html_chunk = tg_format.to_html(chunk)
        try:
            if len(html_chunk) > 4096:
                raise ValueError("too long after formatting")
            await update.message.reply_text(html_chunk, parse_mode="HTML", disable_web_page_preview=True, **extra)
        except Exception:
            for i in range(0, len(chunk), 4096):  # formatting failed: send it plain
                await update.message.reply_text(chunk[i:i + 4096], **(extra if i + 4096 >= len(chunk) else {}))


async def gate(update: Update, count: bool = True):
    """Register the user, enforce blocks and the daily limit. Returns the user row or None."""
    raw, chat_id, name, username = ids(update)
    if getattr(update.effective_chat, "type", "private") != "private":
        chat_id = raw  # a user's private chat id is their user id; never store a group as their address
    uid = store.uid_for("telegram", raw)
    user = store.ensure_user(uid, "telegram", chat_id, name, username)
    if user.get("blocked"):
        return None
    if count and uid != "owner":
        if store.usage_today(uid) >= DAILY_LIMIT:
            await update.message.reply_text(
                f"You've used your {DAILY_LIMIT} requests for today. Come back tomorrow!")
            return None
        store.use_one(uid)
    return user


async def notify_owner(context: ContextTypes.DEFAULT_TYPE, text: str):
    for oid in store.owner_ids()["telegram"]:
        try:
            await context.bot.send_message(chat_id=oid, text=text)
        except Exception:
            pass


async def send_outbox(update: Update, ctx: agent_context.Ctx):
    for item in ctx.outbox:
        try:
            with open(item["path"], "rb") as f:
                if item["kind"] == "document":
                    await update.message.reply_document(document=f, filename=os.path.basename(item["path"]),
                                                        caption=item.get("caption"))
                elif item["kind"] == "video":
                    await update.message.reply_video(
                        video=f, duration=item.get("duration") or None, width=item.get("width") or None,
                        height=item.get("height") or None, supports_streaming=True,
                        read_timeout=300, write_timeout=300)
                else:
                    await update.message.reply_photo(photo=f, read_timeout=120, write_timeout=120)
        except Exception as e:
            await update.message.reply_text(f"Couldn't send the file: {str(e)[:200]}")
        finally:
            if item.get("dir"):
                downloader.cleanup(item)
            elif item["kind"] == "document":
                try:
                    os.unlink(item["path"])
                    os.rmdir(os.path.dirname(item["path"]))
                except OSError:
                    pass
            else:
                try:
                    os.unlink(item["path"])
                except OSError:
                    pass
    ctx.outbox.clear()


# ---------- model choice ----------

PENDING = {}            # picker id -> waiting message
PENDING_TTL = 30 * 60   # buttons expire after 30 minutes
# After picking a model, keep using it until the user is quiet this long (then ask again).
SESSION_HOURS = float(os.getenv("PICKER_SESSION_HOURS", "2"))


def session_model(user: dict) -> str:
    """The model chosen earlier in this conversation, if the user hasn't gone quiet since."""
    key = user.get("session_model") or ""
    if key and time.time() - (user.get("session_at") or 0) < SESSION_HOURS * 3600:
        if key == "auto" or key in models.available():
            return key
    return ""


def _charge_extra(uid: str, key: str) -> bool:
    """Pricier models use more daily requests (the first one was counted on arrival)."""
    if uid == "owner" or key in ("auto", ""):
        return True
    extra = models.weight(key) - 1
    if extra and store.usage_today(uid) + extra > DAILY_LIMIT:
        return False
    for _ in range(extra):
        store.use_one(uid)
    return True


def _model_keyboard(pid: str, is_owner_: bool) -> InlineKeyboardMarkup:
    """Every model, with its version. Ones that aren't connected yet say so."""
    rows = [[InlineKeyboardButton("Auto (picks for you)", callback_data=f"m:{pid}:auto")]]
    buttons = []
    for key, m in models.registry().items():
        label = m["label"]
        if not is_owner_:
            label += f" · {models.weight(key)}"
        if not models.is_ready(key):
            label += " (not set up)"
        buttons.append(InlineKeyboardButton(label, callback_data=f"m:{pid}:{key}"))
    rows += [buttons[i:i + 2] for i in range(0, len(buttons), 2)]
    rows.append([InlineKeyboardButton("Always Auto, stop asking", callback_data=f"m:{pid}:noask")])
    return InlineKeyboardMarkup(rows)


async def run_chat(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str, ask: bool = True):
    """Answer a message. If the user has the picker on and several models exist, ask first."""
    raw, chat_id, name, username = ids(update)
    uid = store.uid_for("telegram", raw)
    user = store.get_user(uid) or {}
    if ask and user.get("picker_on", 1):
        current = session_model(user)
        if current:
            if _charge_extra(uid, current):
                await execute_chat(update, context, text, current)
                return
            store.update_user(uid, session_model="")  # can't afford it any more: ask again
        now = time.time()
        for k in [k for k, v in PENDING.items() if now - v["ts"] > PENDING_TTL]:
            PENDING.pop(k, None)
        pid = secrets.token_urlsafe(6)
        PENDING[pid] = {"text": text, "update": update, "raw": raw, "ts": now}
        note = "" if user.get("uid") == "owner" else f"\n(numbers = requests used; you have {max(0, DAILY_LIMIT - store.usage_today(user['uid']))} left today)"
        await update.message.reply_text("Which model should answer?" + note,
                                        reply_markup=_model_keyboard(pid, user.get("uid") == "owner"))
        return
    key = user.get("default_model") or "auto"
    if not _charge_extra(uid, key):
        key = "auto"
    await execute_chat(update, context, text, key)


async def execute_chat(update, context, text: str, model_key: str, status_msg=None, think: str = None):
    raw, chat_id, name, username = ids(update)
    uid = store.uid_for("telegram", raw)
    if model_key not in ("auto", "") and model_key not in models.available():
        model_key = "auto"
    user = store.get_user(uid) or {}
    if user.get("session_model"):
        store.update_user(uid, session_at=time.time())  # keep the conversation's model alive
    await context.bot.send_chat_action(chat_id=chat_id, action=ChatAction.TYPING)
    ctx = agent_context.Ctx()
    private = getattr(update.effective_chat, "type", "private") == "private"
    live = tg_live.LiveReply(context.bot, chat_id, ctx, private, update.message, rich=uid != "owner").start()
    try:
        reply = await asyncio.to_thread(chat, raw, "telegram", text, ctx, name, username, chat_id, model_key, think)
    finally:
        await live.stop()
    if status_msg is not None:
        try:
            await status_msg.edit_text(f"Answered by {ctx.model_label}")
        except Exception:
            pass
    footer = (f"\n\n_{ctx.model_label} · /switch to change_" if ctx.rich else f"\n\n({ctx.model_label} · /switch to change)") \
        if ctx.model_label and model_key not in ("auto", "") else ""
    await send_long(update, reply + footer, rich=ctx.rich, markup=action_kb(ctx))
    had_files = bool(ctx.outbox)
    await send_outbox(update, ctx)
    if not had_files:
        await offer_download(update.message, text, ctx)
    if uid != "owner" and not store.interests(user) and not store.kv_get(f"topics_offered:{uid}") \
            and getattr(update.effective_chat, "type", "private") == "private":
        await send_topic_picker(update.message, uid,
                                "By the way, want updates on things you follow? " + TOPICS_PROMPT)


async def on_model_pick(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    try:
        _, pid, key = q.data.split(":", 2)
    except ValueError:
        await q.answer()
        return
    pending = PENDING.get(pid)
    if not pending:
        await q.answer("That choice expired. Send your message again.", show_alert=True)
        return
    if str(q.from_user.id) != pending["raw"]:
        await q.answer("This isn't your message.", show_alert=True)
        return
    uid = store.uid_for("telegram", pending["raw"])
    if key == "noask":
        store.update_user(uid, picker_on=0, default_model="auto")
        key = "auto"
    if key != "auto" and key not in models.registry():
        await q.answer("Unknown model.", show_alert=True)
        return
    if key != "auto" and not models.is_ready(key):
        await q.answer(models.setup_hint(key), show_alert=True)
        return
    if not _charge_extra(uid, key):
        await q.answer(f"Not enough requests left today for {models.registry()[key]['short']}. "
                       "Pick Auto or a cheaper model.", show_alert=True)
        return
    store.update_user(uid, session_model=key, session_at=time.time())
    PENDING.pop(pid, None)
    label = "Auto" if key == "auto" else models.registry()[key]["label"]
    await q.answer()
    try:
        await q.edit_message_text(
            f"{label} is thinking..." + ("\n(Stopped asking. Turn it back on with /picker on)" if q.data.endswith(":noask")
                                         else f"\nI'll keep using it for this conversation. /switch to change."))
    except Exception:
        pass
    await execute_chat(pending["update"], context, pending["text"], key, q.message)


async def cmd_model_default(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update, count=False)
    if not user:
        return
    avail = models.available()
    arg = (" ".join(context.args) or "").lower().strip()
    if arg:
        key = next((k for k, m in avail.items() if arg in (k, m["short"].lower())), "auto" if arg == "auto" else None)
        if not key:
            if arg in models.registry():
                await update.message.reply_text(models.setup_hint(arg))
                return
            await update.message.reply_text("Choose one of: auto, " + ", ".join(avail))
            return
        store.update_user(user["uid"], default_model=key)
        await update.message.reply_text(f"Default model: {key}. It's used when the picker is off (/picker off).")
        return
    lines = [f"{k}: {m['label']}" + ("" if user["uid"] == "owner" else f" (uses {models.weight(k)} request{'s' if models.weight(k) > 1 else ''})")
             + ("" if models.is_ready(k) else " - not set up") for k, m in models.registry().items()]
    await update.message.reply_text(
        "Models:\n" + "\n".join(lines) +
        f"\n\nDefault: {user.get('default_model') or 'auto'}. Picker before each message: "
        f"{'on' if user.get('picker_on', 1) else 'off'}.\nSet default: /model sonnet   Toggle picker: /picker on|off")


async def cmd_style(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update, count=False)
    if not user:
        return
    text = " ".join(context.args).strip()
    if text.lower() in ("clear", "reset", "off"):
        store.update_user(user["uid"], style="")
        await update.message.reply_text("Reply style cleared.")
        return
    if not text:
        await update.message.reply_text(
            f"Your reply style: {user.get('style') or 'default'}\n\n"
            "Set one like: /style short and casual, no long intros, mix in Hinglish\nClear it with /style clear")
        return
    store.update_user(user["uid"], style=text[:400])
    await update.message.reply_text(f"Got it. I'll reply like this from now on: {text[:400]}")


async def cmd_thinking(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    from core import agent as _agent
    arg = (" ".join(context.args) or "").lower().strip()
    if arg in _agent.THINKING_LEVELS:
        store.kv_set("thinking_level", arg)
        await update.message.reply_text(
            f"Thinking set to {arg}." + ("" if arg == "off" else
            " Replies get slower and pricier as this goes up (thinking is billed as output). "
            "It applies to your chats only; other users stay on fast replies."))
        return
    await update.message.reply_text(
        f"Thinking is {_agent.thinking_level()}.\n\n"
        "/thinking off - fastest, cheapest\n/thinking low - a quick think on hard questions\n"
        "/thinking medium - balanced\n/thinking high - deepest, slowest, priciest\n\n"
        "Or for one message: /think <question>")


async def cmd_think(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    text = " ".join(context.args).strip()
    if not text:
        await update.message.reply_text("Usage: /think <question>. I'll think hard before answering.")
        return
    user = await gate(update)
    if user:
        await execute_chat(update, context, text, session_model(user) or "sonnet", think="high")


async def cmd_activity(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    arg = " ".join(context.args).strip()
    uid = None
    if arg:
        uid = arg if ":" in arg or arg == "owner" else f"telegram:{arg}"
    rows = store.get_activity(uid=uid, limit=25)
    if not rows:
        await update.message.reply_text("No activity yet." if not uid else f"No activity for {uid}.")
        return
    lines = []
    for r in reversed(rows):
        who = "" if uid else f"{r['uid']} "
        kind = {"message": "asked", "download": "downloaded"}.get(r["kind"], r["kind"].replace("tool:", "used "))
        lines.append(f"{r['ts'][5:16].replace('T', ' ')} {who}{kind}: {r['detail'][:160]}")
    await send_long(update, ("Recent activity" + (f" for {uid}" if uid else "") + " (UTC):\n\n") + "\n".join(lines)
                    + "\n\nFull view: /dashboard")


async def cmd_search(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = " ".join(context.args).strip()
    if not q:
        await update.message.reply_text("Usage: /search <question>. I'll search the web and answer with sources.")
        return
    user = await gate(update)
    if user:
        await run_chat(update, context, (
            f"{q}\n\n(Search the web for this, read the best sources, and answer with numbered citations and a "
            "Sources list at the end.)"))


async def cmd_switch(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update, count=False)
    if not user:
        return
    store.update_user(user["uid"], session_model="", picker_on=1)
    text = " ".join(context.args).strip()
    if text:
        user = await gate(update)
        if user:
            await run_chat(update, context, text)
        return
    await update.message.reply_text("OK. I'll ask which model to use on your next message.")


async def cmd_picker(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update, count=False)
    if not user:
        return
    on = (" ".join(context.args) or "on").lower().startswith("on")
    store.update_user(user["uid"], picker_on=1 if on else 0)
    await update.message.reply_text(
        f"I'll ask which model to use when a conversation starts (after {SESSION_HOURS:g}h of quiet)." if on else
        f"Picker off. I'll use your default model ({user.get('default_model') or 'auto'}); change it with /model.")


# ---------- topic picker ----------

TOPICS_PROMPT = ("What are you interested in? Tap as many as you like, or tap \"Type my own\" and send anything "
                 "(e.g. Kabaddi, K-pop, Tesla). I'll match keywords and send you related updates.")


def _topics_keyboard(uid: str) -> InlineKeyboardMarkup:
    chosen = {t.lower() for t in store.interests(store.get_user(uid) or {})}
    names = list(topics.PRESETS)
    rows, row = [], []
    for i, name in enumerate(names):
        row.append(InlineKeyboardButton(("✓ " if name.lower() in chosen else "") + name, callback_data=f"t:{i}"))
        if len(row) == 2:
            rows.append(row)
            row = []
    if row:
        rows.append(row)
    for i, t in enumerate(store.interests(store.get_user(uid) or {})):
        if not topics.preset_for(t):
            rows.append([InlineKeyboardButton(f"✓ {t} (tap to remove)", callback_data=f"tx:{i}")])
    rows.append([InlineKeyboardButton("Type my own", callback_data="t:own"),
                 InlineKeyboardButton("Done", callback_data="t:done")])
    return InlineKeyboardMarkup(rows)


def _modes_keyboard(current: str = "") -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup([[InlineKeyboardButton(("✓ " if k == current else "") + label, callback_data=f"f:{k}")]
                                 for k, label in topics.MODES.items()])


def _tz_keyboard() -> InlineKeyboardMarkup:
    rows, row = [], []
    for i, (label, _) in enumerate(topics.TIMEZONES):
        row.append(InlineKeyboardButton(label, callback_data=f"z:{i}"))
        if len(row) == 2:
            rows.append(row)
            row = []
    if row:
        rows.append(row)
    rows.append([InlineKeyboardButton("Somewhere else (type your city)", callback_data="z:own")])
    return InlineKeyboardMarkup(rows)


def _summary(u: dict) -> str:
    mode = u.get("news_mode") or "both"
    if not u.get("news_on"):
        mode = "off"
    when = {"live": "as stories break (7 AM to 11 PM, max 8 a day) plus a brief at 8 AM",
            "both": "at 8 AM and 7 PM", "morning": "at 8 AM", "off": ""}[mode]
    ints = ", ".join(store.interests(u)) or "nothing yet"
    if mode == "off":
        return f"Topics: {ints}. Updates are off; turn them on with /updates."
    return f"Topics: {ints}.\nUpdates {when}, your time ({u.get('timezone')}).\nChange anytime: /interests, /updates"


async def send_topic_picker(message, uid: str, intro: str = TOPICS_PROMPT):
    store.kv_set(f"topics_offered:{uid}", True)
    await message.reply_text(intro, reply_markup=_topics_keyboard(uid))


async def _finish_setup(q, uid: str):
    u = store.get_user(uid) or {}
    await q.edit_message_text("All set.\n\n" + _summary(u))
    if u.get("news_on") and store.interests(u):
        preview = await asyncio.to_thread(topics.latest_now, uid)
        if preview:
            try:
                await q.message.reply_text(preview, parse_mode="HTML", disable_web_page_preview=True,
                                           reply_markup=links_kb(uid))
            except Exception:
                pass


async def on_topic_pick(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    uid = store.uid_for("telegram", str(update.effective_user.id))
    user = store.get_user(uid)
    if not user or user.get("blocked"):
        await q.answer()
        return
    data = q.data
    kind, _, val = data.partition(":")
    if kind == "t" and val.isdigit():
        names = list(topics.PRESETS)
        if int(val) < len(names):
            ints = store.interests(user)
            name = names[int(val)]
            if name.lower() not in [i.lower() for i in ints] and len(ints) >= 12:
                await q.answer("You can follow up to 12 topics. Remove one first.", show_alert=True)
                return
            topics.toggle(uid, name)
        await q.answer()
        await q.edit_message_reply_markup(reply_markup=_topics_keyboard(uid))
    elif kind == "tx":
        ints = store.interests(user)
        if val.isdigit() and int(val) < len(ints):
            topics.toggle(uid, ints[int(val)])
        await q.answer("Removed")
        await q.edit_message_reply_markup(reply_markup=_topics_keyboard(uid))
    elif data == "t:own":
        store.update_user(uid, state="awaiting_topics")
        await q.answer()
        await q.message.reply_text("Send your topics in one message, separated by commas. For example: "
                                   "Kabaddi, K-pop, Tesla, Chennai Super Kings")
    elif data == "t:done":
        await q.answer()
        if not store.interests(user):
            await q.edit_message_text("No topics picked, so no updates for now. Pick some anytime with /interests.")
            return
        await q.edit_message_text("Topics: " + ", ".join(store.interests(user)) + "\n\nHow often do you want updates?",
                                  reply_markup=_modes_keyboard(user.get("news_mode") or ""))
    elif kind == "f" and val in topics.MODES:
        store.update_user(uid, news_mode=val, news_on=0 if val == "off" else 1)
        await q.answer()
        if val != "off" and not store.kv_get(f"tz_chosen:{uid}"):
            await q.edit_message_text("Where are you? So updates arrive at the right local time.",
                                      reply_markup=_tz_keyboard())
        else:
            await _finish_setup(q, uid)
    elif kind == "z":
        await q.answer()
        if val == "own":
            store.update_user(uid, state="awaiting_city")
            await q.edit_message_text("Send me your city (for example: Toronto, Berlin, Nairobi).")
            return
        if val.isdigit() and int(val) < len(topics.TIMEZONES):
            store.update_user(uid, timezone=topics.TIMEZONES[int(val)][1])
            store.kv_set(f"tz_chosen:{uid}", True)
        await _finish_setup(q, uid)
    else:
        await q.answer()


async def receive_typed_topics(update: Update, user: dict, text: str) -> bool:
    """Handles the message after 'Type my own' / 'type your city'. Returns True if consumed."""
    state = user.get("state")
    if state not in ("awaiting_topics", "awaiting_city"):
        return False
    uid = user["uid"]
    store.update_user(uid, state="active")
    if "?" in text or len(text) > 120:
        return False  # looks like a question, not a topic list: just answer it
    if state == "awaiting_city":
        try:
            info = await asyncio.to_thread(profile.extract_onboarding, f"I'm in {text}", uid)
        except Exception:
            info = {}
        if info.get("timezone"):
            store.update_user(uid, timezone=info["timezone"])
            store.kv_set(f"tz_chosen:{uid}", True)
            store.add_fact(uid, f"Based in {text.strip()[:60]}")
        msg = await update.message.reply_text("Saved." if info.get("timezone") else
                                              "I couldn't place that city, so I'll keep " + (user.get("timezone") or "")
                                              + ". Try /interests again anytime.")
        u = store.get_user(uid) or {}
        await update.message.reply_text(_summary(u))
        if u.get("news_on") and store.interests(u):
            preview = await asyncio.to_thread(topics.latest_now, uid)
            if preview:
                await update.message.reply_text(preview, parse_mode="HTML", disable_web_page_preview=True,
                                                reply_markup=links_kb(uid))
        return True
    ints = topics.add_custom(uid, text)
    await update.message.reply_text("Added. Tap more or press Done.\n\nYour topics: " + ", ".join(ints),
                                    reply_markup=_topics_keyboard(uid))
    return True


async def cmd_updates(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update, count=False)
    if not user:
        return
    current = (user.get("news_mode") or "both") if user.get("news_on") else "off"
    await update.message.reply_text(_summary(user) + "\n\nHow often do you want updates?",
                                    reply_markup=_modes_keyboard(current))


# ---------- onboarding ----------

async def start_onboarding(update: Update, context: ContextTypes.DEFAULT_TYPE, user: dict):
    store.update_user(user["uid"], state="active")
    first = update.effective_user.first_name
    await update.message.reply_text(WELCOME.format(name=f" {first}" if first else ""))
    await send_topic_picker(update.message, user["uid"])
    if user.get("state") == "new":
        uname = f"@{user['username']}" if user.get("username") else user["uid"]
        await notify_owner(context, f"New user started the bot: {first or ''} {uname} ({user['uid']})")


async def finish_onboarding(update: Update, user: dict, text: str):
    try:
        info = await asyncio.to_thread(profile.extract_onboarding, text, user["uid"])
    except Exception as e:
        await update.message.reply_text(f"Sorry, something went wrong reading that ({str(e)[:100]}). Try once more?")
        return
    if not info["interests"]:
        await update.message.reply_text(
            "I didn't catch any topics. What do you like to follow? For example: cricket, AI, startups, fashion, stock markets.")
        return
    await asyncio.to_thread(profile.apply_onboarding, user["uid"], info)
    topics.attach_presets(user["uid"])
    if info.get("timezone"):
        store.kv_set(f"tz_chosen:{user['uid']}", True)
    u = store.get_user(user["uid"]) or {}
    await update.message.reply_text("Got it. I'll follow: " + ", ".join(store.interests(u))
                                    + ".\n\nHow often do you want updates?",
                                    reply_markup=_modes_keyboard(u.get("news_mode") or ""))


# ---------- commands ----------

async def cmd_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update, count=False)
    if not user:
        return
    if user["uid"] == "owner":
        await update.message.reply_text("Agent online.\n\n" + _bot_name(OWNER_HELP, context))
        return
    if user.get("state") == "active":
        await update.message.reply_text("Welcome back! " + _bot_name(HELP, context) + PRIVACY)
        return
    await start_onboarding(update, context, user)


async def cmd_help(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update, count=False)
    if user:
        await update.message.reply_text(_bot_name(OWNER_HELP if user["uid"] == "owner" else HELP + PRIVACY, context))


async def cmd_clear(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update, count=False)
    if user:
        memory.clear(user["uid"], "telegram")
        store.update_user(user["uid"], session_model="")
        store.kv_set(f"pfile:{user['uid']}", None)
        await update.message.reply_text("New conversation started. (What I've learned about you is kept; /forget erases it.)")


async def cmd_profile(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update, count=False)
    if not user:
        return
    block = profile.profile_block(user["uid"])
    news = "on" if user.get("news_on") else "off"
    await send_long(update, (block.split("\n", 1)[1] if block else "I don't know anything about you yet.")
                    + f"\n\nDaily news: {news}. Erase everything with /forget.")


async def cmd_forget(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update, count=False)
    if not user:
        return
    store.forget_user(user["uid"])
    await update.message.reply_text(
        "Done. I've erased your profile, interests, saved facts, notes, reminders and our conversation."
        + ("" if user["uid"] == "owner" else " Send /start to set up again."))


async def cmd_interests(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update)
    if not user:
        return
    text = " ".join(context.args)
    if not text:
        await send_topic_picker(update.message, user["uid"])
        return
    await finish_onboarding(update, user, text)


async def cmd_keywords(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update, count=False)
    if not user:
        return
    arg = " ".join(context.args).strip()
    if arg.lower() in ("clear", "reset"):
        store.update_user(user["uid"], custom_keywords=[])
        await update.message.reply_text("Your custom keywords are cleared.")
        return
    if arg:
        new = [k.strip()[:50] for k in arg.split(",") if k.strip()]
        merged = list(dict.fromkeys(store.custom_keywords(user) + new))[:15]
        store.update_user(user["uid"], custom_keywords=merged)
        await update.message.reply_text(
            "Tracking: " + ", ".join(merged) + "\n\nI'll match these across news sites, Reddit, X, Instagram, "
            "Threads, YouTube and Hacker News in every brief. Clear them with /keywords clear.")
        return
    kw = await asyncio.to_thread(profile.ensure_keywords, user["uid"])
    lines = [f"{t}: {', '.join(k) or '(topic name only)'}" for t, k in kw.items()]
    custom = store.custom_keywords(user)
    if custom:
        lines.append("Your own keywords: " + ", ".join(custom))
    await update.message.reply_text(
        ("Keywords I match in your news:\n\n" + "\n".join(lines) if lines else "No topics yet. Use /interests first.")
        + "\n\nAdd your own: /keywords Zepto, IPL auction, Nothing Phone")


async def cmd_alert(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update, count=False)
    if not user:
        return
    from core import watch
    kw = " ".join(context.args)
    if not kw:
        await update.message.reply_text("Tell me what to watch, like: /alert Zepto\n\n" + watch.list_text(user["uid"]))
        return
    await update.message.reply_text(watch.add(user["uid"], str(update.effective_chat.id), kw))


async def cmd_alerts(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update, count=False)
    if user:
        from core import watch
        await update.message.reply_text(watch.list_text(user["uid"]))


async def cmd_unalert(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update, count=False)
    if not user:
        return
    from core import watch
    what = " ".join(context.args)
    await update.message.reply_text(watch.remove(user["uid"], what) if what else "Usage: /unalert <id or keyword>")


async def cmd_pipeline(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    from core import talent
    await send_long(update, talent.pipeline_text())


def _downtime_line() -> str:
    week = [d for d in store.kv_get("downtime_log", []) if d["to"] > time.time() - 7 * 86400]
    if not week:
        return "Offline in the last 7 days: none recorded"
    mins = int(sum(d["to"] - d["from"] for d in week) // 60)
    return f"Offline in the last 7 days: {len(week)} times, {mins // 60}h {mins % 60}m total"


def _ago(ts) -> str:
    import time as _t
    if not ts:
        return "never"
    s = int(_t.time() - ts)
    return f"{s // 3600}h {s % 3600 // 60}m ago" if s >= 3600 else (f"{s // 60}m ago" if s >= 60 else f"{s}s ago")


async def cmd_status(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    import time as _t
    from core import sources, watch
    up = int(_t.time() - alerts.STARTED)
    tick = store.kv_get("scheduler_tick")
    lines = [
        f"Running for {up // 3600}h {up % 3600 // 60}m",
        f"Scheduler: last ran {_ago(tick)}" + ("  << STUCK?" if tick and _t.time() - tick > 300 else ""),
        _downtime_line(),
        "",
        f"Spend today: ${store.spend_today():.3f} (you ${store.spend_today(owner=True):.3f}, others "
        f"${store.spend_today(owner=False):.3f} of ${llm.PUBLIC_DAILY_BUDGET:.2f} budget)",
        f"Serper searches today: {store.counter('serper')}",
        (f"Anthropic credits left (estimate): ${store.credit_estimate()['left']:.2f}" if store.credit_estimate()
         else "Anthropic credits: unknown (send /credits <balance>)"),
        "Models: " + ", ".join(m["short"] for m in models.available().values()),
        "",
    ]
    users = store.all_users()
    active = [u for u in users if (u.get("last_seen") or "").startswith(store.iso(store.now_utc())[:10])]
    lines.append(f"Users: {len(users)} total, {len(active)} active today, "
                 f"{sum(1 for u in users if u.get('news_on'))} getting news")
    lines.append(f"Keyword alerts: {len(store.list_watches())} active (every {watch.ALERT_EVERY_MIN} min)")
    lines.append(f"Voice notes: {'ready' if voice.available() else 'not installed'}")
    lines.append("\nSources:")
    if not sources.HEALTH:
        lines.append("  no checks yet (send /news to run them)")
    for kind, h in sources.HEALTH.items():
        name = sources.SOURCE_NAMES.get(kind, kind)
        if h.get("fails"):
            lines.append(f"  {name}: FAILING x{h['fails']} ({h.get('err', '')[:80]})")
        else:
            lines.append(f"  {name}: ok, {_ago(h.get('ok'))}")
    await send_long(update, "\n".join(lines))


async def cmd_usage(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    out = []
    for label, days in (("TODAY", 1), ("LAST 7 DAYS", 7)):
        b = store.spend_breakdown(days)
        out.append(f"{label}: ${b['total']:.3f}, Serper searches: {store.counter('serper', days)}")
        out += [f"  {f['feature']}: ${f['cost']:.3f} ({f['calls']} calls)" for f in b["by_feature"]]
        if days == 7 and b["by_user"]:
            out.append("  Top users:")
            out += [f"    {u['uid']}: ${u['cost']:.3f}" for u in b["by_user"][:5]]
        out.append("")
    week = store.spend_breakdown(7)["total"]
    out.append(f"At this rate: about ${week / 7 * 30:.2f} a month.")
    await send_long(update, "\n".join(out))


async def cmd_credits(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    arg = " ".join(context.args).replace("$", "").strip()
    if arg:
        try:
            amount = float(arg)
        except ValueError:
            await update.message.reply_text("Usage: /credits 6.54 (your current balance at console.anthropic.com > Billing)")
            return
        store.kv_set("credit_baseline", {"amount": amount, "set_at": store.iso(store.now_utc())})
        await update.message.reply_text(f"Got it: ${amount:.2f}. I'll subtract Claude spend from here on. "
                                        "Update it whenever you top up.")
        return
    est = store.credit_estimate()
    await update.message.reply_text(
        f"Estimated Anthropic credits left: ${est['left']:.2f} (from ${est['entered']:.2f} entered)." if est else
        "Tell me your balance first: /credits 6.54 (from console.anthropic.com > Billing).")


async def cmd_dashboard(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    from platforms import dashboard
    await update.message.reply_text(
        f"Dashboard: http://localhost:{dashboard.PORT}\nOpen it in a browser on the computer running the agent. "
        "It shows active sessions, each user's requests and spend, spend by model and day, and source health.")


async def on_error(update, context: ContextTypes.DEFAULT_TYPE):
    err = context.error
    print(f"Telegram handler error: {err!r}")
    alerts.notify(f"tg:{type(err).__name__}", f"A Telegram request crashed: {type(err).__name__}: {str(err)[:200]}", 60)
    try:
        if update and getattr(update, "message", None):
            await update.message.reply_text("Sorry, something went wrong on my side. Try again in a moment.")
    except Exception:
        pass


async def cmd_news_toggle(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update, count=False)
    if not user:
        return
    on = update.message.text.lower().startswith("/newson")
    store.update_user(user["uid"], news_on=1 if on else 0)
    await update.message.reply_text("Daily news is on." if on else "Daily news is off. Turn it back on with /newson.")


async def cmd_news(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update)
    if not user:
        return
    from core.news import build_user_brief
    await update.message.reply_text("Pulling your news...")
    try:
        brief = await asyncio.to_thread(build_user_brief, user["uid"])
        kb = None if brief.startswith("Nothing new") else links_kb(user["uid"])
    except Exception as e:
        brief, kb = f"Couldn't build your news: {str(e)[:200]}", None
    await send_long(update, brief, markup=kb)


async def cmd_brief(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return await cmd_news(update, context)
    await update.message.reply_text("Building your briefing. Give me 2-3 minutes...")
    from digest import build_digest
    try:
        brief = await asyncio.to_thread(build_digest)
    except Exception as e:
        brief = f"Could not build the briefing: {str(e)[:200]}"
    await send_long(update, brief)


async def cmd_radar(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    await update.message.reply_text("Scanning news, Reddit, X, Instagram and YouTube across your radar topics. About a minute...")
    from core.radar import build_radar
    try:
        text = await asyncio.to_thread(build_radar)
    except Exception as e:
        text = f"Radar failed: {str(e)[:200]}"
    await send_long(update, text)


async def cmd_users(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    users = [u for u in store.all_users() if u["uid"] != "owner"]
    if not users:
        await update.message.reply_text("No other users yet.")
        return
    lines = []
    for u in users:
        tag = f"@{u['username']}" if u.get("username") else (u.get("name") or "")
        flags = " BLOCKED" if u.get("blocked") else ""
        lines.append(f"{u['uid'].split(':')[-1]} {tag}: {store.usage_today(u['uid'])}/{DAILY_LIMIT} today, "
                     f"{u.get('msg_count') or 0} msgs total, news {'on' if u.get('news_on') else 'off'}{flags}")
    await send_long(update, f"{len(users)} users\n\n" + "\n".join(lines))


async def cmd_block(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update) or not context.args:
        return
    target = f"telegram:{context.args[0].strip()}"
    if not store.get_user(target):
        await update.message.reply_text("No user with that ID. See /users.")
        return
    block = update.message.text.lower().startswith("/block")
    store.update_user(target, blocked=1 if block else 0)
    await update.message.reply_text(("Blocked " if block else "Unblocked ") + target)


async def cmd_model(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """/deep and /quick: skip the picker and pass the whole text (with the prefix) to the agent."""
    user = await gate(update)
    if user:
        await execute_chat(update, context, update.message.text, "auto")


# ---------- messages ----------

async def process_text(update: Update, context: ContextTypes.DEFAULT_TYPE, user: dict, text: str):
    if user["uid"] != "owner" and user.get("state") in ("new", "awaiting_interests", None, ""):
        # Like ChatGPT: answer straight away. Daily news is optional (/interests).
        store.update_user(user["uid"], state="active")
        if user.get("state") == "new":
            first = update.effective_user.first_name
            await update.message.reply_text(FIRST_HELLO.format(name=f" {first}" if first else ""))
            uname = f"@{user['username']}" if user.get("username") else user["uid"]
            await notify_owner(context, f"New user: {first or ''} {uname} ({user['uid']})")

    if await receive_typed_topics(update, user, text):
        return
    waiting = pending_file(user["uid"])
    if waiting and not text.startswith("/"):
        await run_file_action(update, context, user, waiting, text, "typed")
        return
    if downloader.wants_download(text):
        await download_link(update, context, downloader.find_url(text))
        return
    await run_chat(update, context, link_prompt(text))


def link_prompt(text: str) -> str:
    """A message that's only a link becomes 'tell me about this link' for the agent."""
    if downloader.only_link(text):
        return f"{downloader.any_url(text)}\n\n{downloader.BARE_LINK_NOTE}"
    return text


async def offer_download(message, text: str, ctx):
    """After talking about a bare video link, offer the file with one tap."""
    if downloader.BARE_LINK_NOTE not in text or ctx.outbox:
        return
    url = downloader.any_url(text)
    if not downloader.is_video_link(url):
        return
    key = secrets.token_urlsafe(6)
    store.kv_set(f"dlurl:{key}", url)
    try:
        await message.reply_text("Want the video file?", reply_markup=InlineKeyboardMarkup(
            [[InlineKeyboardButton("Download video", callback_data=f"dl:{key}")]]))
    except Exception:
        pass


async def on_download_pick(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    url = store.kv_get(f"dlurl:{q.data.split(':', 1)[1]}")
    await q.answer()
    if not url:
        await q.edit_message_text("That button expired. Send the link again with the word download.")
        return
    user = store.get_user(store.uid_for("telegram", str(update.effective_user.id)))
    if user and user.get("blocked"):
        return
    try:
        await q.edit_message_reply_markup(reply_markup=None)
    except Exception:
        pass
    shim = types.SimpleNamespace(message=q.message, effective_user=update.effective_user,
                                 effective_chat=update.effective_chat)
    await download_link(shim, context, url)


async def download_link(update: Update, context: ContextTypes.DEFAULT_TYPE, url: str):
    store.log_activity(store.uid_for("telegram", str(update.effective_user.id)), "telegram", "download", url)
    status = await update.message.reply_text("Downloading...")
    await context.bot.send_chat_action(chat_id=update.effective_chat.id, action=ChatAction.UPLOAD_VIDEO)
    try:
        result = await asyncio.to_thread(downloader.download_video, url)
    except downloader.DownloadFailed as e:
        await status.edit_text(str(e))
        return
    except Exception as e:
        await status.edit_text(f"Something went wrong: {str(e)[:200]}")
        return
    await status.edit_text(f"Uploading {result['size_mb']} MB...")
    ctx = agent_context.Ctx(outbox=[{"kind": "video", **result}])
    await send_outbox(update, ctx)
    try:
        await status.delete()
    except Exception:
        pass


async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update)
    if user:
        await process_text(update, context, user, update.message.text)


async def handle_voice(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = await gate(update)
    if not user:
        return
    v = update.message.voice or update.message.audio
    if not voice.available():
        await update.message.reply_text("Voice notes aren't set up on this server yet (pip install faster-whisper).")
        return
    if v.duration and v.duration > voice.MAX_SECONDS:
        await update.message.reply_text(f"That's a bit long. Send voice notes under {voice.MAX_SECONDS // 60} minutes.")
        return
    await context.bot.send_chat_action(chat_id=update.effective_chat.id, action=ChatAction.TYPING)
    f = await context.bot.get_file(v.file_id)
    with tempfile.NamedTemporaryFile(delete=False, suffix=".ogg") as tmp:
        await f.download_to_drive(tmp.name)
        path = tmp.name
    try:
        text, lang = await asyncio.to_thread(voice.transcribe, path)
    except Exception as e:
        await update.message.reply_text(f"Couldn't transcribe that: {str(e)[:200]}")
        return
    finally:
        os.unlink(path)
    if not text:
        await update.message.reply_text("I couldn't hear anything in that voice note.")
        return
    await process_text(update, context, user, text)


# ---------- files and images ----------

def extract_pdf(path: str) -> str:
    return "".join(page.get_text() for page in fitz.open(path))[:8000]


def extract_pptx(path: str) -> str:
    out = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        out.append(f"Slide {i}: " + " ".join(s.text for s in slide.shapes if hasattr(s, "text")))
    return "\n".join(out)[:8000]


def extract_docx(path: str) -> str:
    return "\n".join(p.text for p in Document(path).paragraphs)[:8000]


def extract_txt(path: str) -> str:
    with open(path, "r", errors="ignore") as f:
        return f.read()[:8000]


def extract_excel(path: str) -> str:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    text = ""
    for ws in wb.worksheets:
        text += f"Sheet: {ws.title}\n"
        for row in ws.iter_rows(values_only=True):
            cells = " | ".join(str(c) for c in row if c is not None)
            if cells.strip():
                text += cells + "\n"
            if len(text) > 8000:
                break
    return text[:8000]


EXTRACTORS = {
    ".pdf": extract_pdf, ".pptx": extract_pptx, ".docx": extract_docx,
    ".txt": extract_txt, ".csv": extract_txt, ".md": extract_txt, ".xlsx": extract_excel,
}


def describe_image(path: str, owner: bool, uid: str = None) -> str:
    with open(path, "rb") as f:
        b64 = base64.standard_b64encode(f.read()).decode()
    ext = path.rsplit(".", 1)[-1].lower()
    media = {"png": "image/png", "gif": "image/gif", "webp": "image/webp"}.get(ext, "image/jpeg")
    resp = client.messages.create(
        _uid=uid,
        model=VISION_OWNER_MODEL if owner else VISION_PUBLIC_MODEL,
        max_tokens=900,
        messages=[{"role": "user", "content": [
            {"type": "image", "source": {"type": "base64", "media_type": media, "data": b64}},
            {"type": "text", "text": "Describe this image in detail. Extract any text exactly. "
                                     "Plain text, no markdown."},
        ]}],
    )
    return resp.content[0].text.strip()


async def handle_image_file(update: Update, context: ContextTypes.DEFAULT_TYPE, user: dict, path: str):
    await context.bot.send_chat_action(chat_id=update.effective_chat.id, action=ChatAction.TYPING)
    try:
        desc = await asyncio.to_thread(describe_image, path, user["uid"] == "owner", user["uid"])
    except Exception as e:
        await update.message.reply_text(f"Could not read image: {str(e)[:200]}")
        return
    caption = (update.message.caption or "").strip()
    if caption:
        # Answer the question about the image, with the description in memory for follow-ups.
        await run_chat(update, context, f"[I sent an image. What it shows: {desc}]\n\n{caption}")
    else:
        memory.add_exchange(user["uid"], "telegram", f"[I sent an image. What it shows: {desc}]", desc)
        await send_long(update, desc)


FILE_ACTIONS = {
    "doc": [("sum", "Summarise", "Summarise this file: the main point first, then the important details."),
            ("key", "Key points", "List the key points and takeaways as short bullets."),
            ("ana", "Analyse", "Analyse it: main insights, strengths, weaknesses, risks, and what to do next."),
            ("ask", "Ask something", None)],
    "image": [("desc", "Describe", None),
              ("text", "Read the text", "Give me all the text in this image exactly, keeping its layout where it helps."),
              ("ana", "Analyse", "Analyse this image: what it shows, what stands out, and anything important or wrong."),
              ("ask", "Ask something", None)],
}
PENDING_FILE_MIN = 30


def _file_kb(kind: str, key: str):
    acts = FILE_ACTIONS[kind]
    rows = [[InlineKeyboardButton(label, callback_data=f"fa:{key}:{code}") for code, label, _ in acts[i:i + 2]]
            for i in range(0, len(acts), 2)]
    return InlineKeyboardMarkup(rows)


async def offer_file_actions(update: Update, uid: str, file_id: str, name: str, ext: str, kind: str):
    key = secrets.token_urlsafe(5)
    store.kv_set(f"pfile:{uid}", {"key": key, "file_id": file_id, "name": name, "ext": ext, "kind": kind,
                                 "at": time.time()})
    what = "this image" if kind == "image" else f"{name}"
    await update.message.reply_text(
        f"Got {what}. What should I do with it?\nTap an option, or just type what you want "
        "(for example: \"make notes for my exam\", \"find the numbers\", \"translate to Hindi\").",
        reply_markup=_file_kb(kind, key))


def pending_file(uid: str):
    p = store.kv_get(f"pfile:{uid}")
    if p and time.time() - p.get("at", 0) < PENDING_FILE_MIN * 60:
        return p
    return None


async def run_file_action(update, context, user: dict, p: dict, instruction: str = None, action: str = None):
    """Download the waiting file, read it, and do what was asked (counts as one request)."""
    uid = user["uid"]
    store.kv_set(f"pfile:{uid}", None)
    if uid != "owner":
        if store.usage_today(uid) >= DAILY_LIMIT:
            await update.message.reply_text(f"You've used your {DAILY_LIMIT} requests for today. Come back tomorrow!")
            return
        store.use_one(uid)
    f = await context.bot.get_file(p["file_id"])
    with tempfile.NamedTemporaryFile(delete=False, suffix=p["ext"]) as tmp:
        await f.download_to_drive(tmp.name)
        path = tmp.name
    try:
        if p["kind"] == "image":
            await context.bot.send_chat_action(chat_id=update.effective_chat.id, action=ChatAction.TYPING)
            desc = await asyncio.to_thread(describe_image, path, uid == "owner", uid)
            if action == "desc":
                memory.add_exchange(uid, "telegram", "[I sent an image]", desc)
                await send_long(update, desc)
                return
            await run_chat(update, context, f"[I sent an image. What it shows: {desc}]\n\n{instruction}", ask=False)
            return
        await update.message.reply_text(f"Reading {p['name']}...")
        text = await asyncio.to_thread(EXTRACTORS[p["ext"]], path)
        await run_chat(update, context, f"I'm sending you a file called '{p['name']}'. {instruction}\n\nFile content:\n{text}",
                       ask=False)
    except Exception as e:
        await update.message.reply_text(f"Could not read the file: {str(e)[:200]}")
    finally:
        os.unlink(path)


async def on_file_action(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    try:
        _, key, code = q.data.split(":", 2)
    except ValueError:
        return
    uid = store.uid_for("telegram", str(update.effective_user.id))
    user = store.get_user(uid)
    p = pending_file(uid)
    if not user or user.get("blocked"):
        return
    if not p or p["key"] != key:
        await q.message.reply_text("That file has expired. Send it again.")
        return
    try:
        await q.edit_message_reply_markup(reply_markup=None)
    except Exception:
        pass
    shim = types.SimpleNamespace(message=q.message, effective_user=update.effective_user,
                                 effective_chat=update.effective_chat)
    if code == "ask":
        store.kv_set(f"pfile:{uid}", {**p, "at": time.time()})
        await q.message.reply_text("Type your question about " + ("the image." if p["kind"] == "image" else f"{p['name']}."))
        return
    label, instruction = next(((lb, ins) for c, lb, ins in FILE_ACTIONS[p["kind"]] if c == code), (None, None))
    if label is None:
        return
    await run_file_action(shim, context, user, p, instruction, code)


async def handle_photo(update: Update, context: ContextTypes.DEFAULT_TYPE):
    caption = (update.message.caption or "").strip()
    user = await gate(update, count=bool(caption))
    if not user:
        return
    if not caption:
        await offer_file_actions(update, user["uid"], update.message.photo[-1].file_id, "photo", ".jpg", "image")
        return
    f = await context.bot.get_file(update.message.photo[-1].file_id)
    with tempfile.NamedTemporaryFile(delete=False, suffix=".jpg") as tmp:
        await f.download_to_drive(tmp.name)
        path = tmp.name
    try:
        await handle_image_file(update, context, user, path)
    finally:
        os.unlink(path)


async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    caption = (update.message.caption or "").strip()
    user = await gate(update, count=bool(caption))
    if not user:
        return
    doc = update.message.document
    name = doc.file_name or "file"
    ext = os.path.splitext(name.lower())[1]
    if doc.file_size and doc.file_size > 20 * 1024 * 1024:
        await update.message.reply_text("That file is over 20 MB, which is Telegram's limit for bots.")
        return
    is_image = ext in (".jpg", ".jpeg", ".png", ".gif", ".webp")
    if not is_image and ext not in EXTRACTORS:
        await update.message.reply_text(
            "I can read PDF, Word (.docx), PowerPoint (.pptx), Excel (.xlsx), CSV, TXT and images. "
            "For old .doc/.ppt/.xls files, save them in the new format first.")
        return
    if not caption:  # ask what they want before spending anything
        await offer_file_actions(update, user["uid"], doc.file_id, name, ext, "image" if is_image else "doc")
        return
    f = await context.bot.get_file(doc.file_id)
    with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
        await f.download_to_drive(tmp.name)
        path = tmp.name
    try:
        if is_image:
            await handle_image_file(update, context, user, path)
            return
        await update.message.reply_text("Reading your file...")
        try:
            text = await asyncio.to_thread(EXTRACTORS[ext], path)
        except Exception as e:
            await update.message.reply_text(f"Could not read file: {str(e)[:200]}")
            return
        ask = (update.message.caption or "").strip() or \
            "Read it and summarise the key points, insights, and anything important."
        await run_chat(update, context, f"I'm sending you a file called '{name}'. {ask}\n\nFile content:\n{text}")
    finally:
        os.unlink(path)


# ---------- group chats (Meta AI style: mention me or reply to me) ----------

GROUPS_ENABLED = os.getenv("GROUPS_ENABLED", "true").lower() == "true"
GROUP_INTRO = ("Hi everyone! I'm an AI assistant built by Naveen Raj. Mention me (@{bot}) or reply to one of my "
               "messages and I'll answer: questions, research, news, links, summaries, video downloads.\n"
               "Admins: /botoff pauses me here, /boton brings me back.\n"
               "Messages sent to me are saved on the bot owner's hard disk, and the owner can read them.")


def _addressed_to_bot(msg, bot) -> tuple:
    """(is_for_me, text_without_mention)."""
    text = msg.text or msg.caption or ""
    entities = (msg.entities or []) + (msg.caption_entities or [])
    mentioned = False
    for e in entities:
        if e.type == "mention" and text[e.offset:e.offset + e.length].lower() == f"@{bot.username}".lower():
            mentioned = True
        elif e.type == "text_mention" and e.user and e.user.id == bot.id:
            mentioned = True
    replied = bool(msg.reply_to_message and msg.reply_to_message.from_user
                   and msg.reply_to_message.from_user.id == bot.id)
    clean = re.sub(rf"@{re.escape(bot.username)}\b", "", text, flags=re.I).strip() if bot.username else text
    return (mentioned or replied), clean


async def handle_group(update: Update, context: ContextTypes.DEFAULT_TYPE):
    msg = update.message
    if not GROUPS_ENABLED or not msg:
        return
    for_me, text = _addressed_to_bot(msg, context.bot)
    if not for_me:
        return
    sender = getattr(msg, "from_user", None) or update.effective_user
    if sender is None or getattr(sender, "is_bot", False):
        return  # never answer other bots or anonymous channel posts (avoids bot-to-bot loops)
    if tg_live.group_off(update.effective_chat.id):
        return
    if not text:
        await msg.reply_text("Yes? Ask me anything: mention me with your question, or reply to my message.")
        return
    ok, warning = tg_live.check_group(update.effective_chat.id, sender.id)
    if not ok:
        if warning:
            await msg.reply_text(warning)
        return
    text = text[:tg_live.GROUP_MAX_CHARS]
    user = await gate(update)
    if not user:
        return
    if downloader.wants_download(text):
        await download_link(update, context, downloader.find_url(text))
        return
    text = link_prompt(text)
    quoted = msg.reply_to_message
    if quoted and quoted.from_user and quoted.from_user.id != context.bot.id and (quoted.text or quoted.caption):
        who = quoted.from_user.first_name or "someone"
        text = f"[Replying to {who}'s message: \"{(quoted.text or quoted.caption)[:1500]}\"]\n{text}"
    raw, _, name, username = ids(update)
    chat = update.effective_chat
    await context.bot.send_chat_action(chat_id=chat.id, action=ChatAction.TYPING)
    ctx = agent_context.Ctx()
    live = tg_live.LiveReply(context.bot, chat.id, ctx, False, msg).start()
    try:
        reply = await asyncio.to_thread(chat_fn, raw, "telegram", text, ctx, name, username, raw, "auto", None,
                                        {"id": chat.id, "title": chat.title})
    finally:
        await live.stop()
    await send_long(update, reply, rich=ctx.rich, markup=action_kb(ctx))
    had_files = bool(ctx.outbox)
    await send_outbox(update, ctx)
    if not had_files:
        await offer_download(msg, text, ctx)


async def cmd_help_group(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(GROUP_INTRO.format(bot=context.bot.username) + "\n\nOr use /ask <question>.")


async def cmd_ask(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """/ask <question> works in groups too (handy when the bot's privacy mode hides normal messages)."""
    text = " ".join(context.args).strip()
    if not text:
        await update.message.reply_text("Usage: /ask <your question>")
        return
    chat = update.effective_chat
    in_group = getattr(chat, "type", "private") != "private"
    if in_group:
        if tg_live.group_off(chat.id) or (update.effective_user and update.effective_user.is_bot):
            return
        ok, warning = tg_live.check_group(chat.id, update.effective_user.id)
        if not ok:
            if warning:
                await update.message.reply_text(warning)
            return
        text = text[:tg_live.GROUP_MAX_CHARS]
    user = await gate(update)
    if not user:
        return
    raw, _, name, username = ids(update)
    group = {"id": chat.id, "title": chat.title} if in_group else None
    ctx = agent_context.Ctx()
    live = tg_live.LiveReply(context.bot, chat.id, ctx, not in_group, update.message,
                             rich=user["uid"] != "owner").start()
    try:
        reply = await asyncio.to_thread(chat_fn, raw, "telegram", text, ctx, name, username, raw, "auto", None, group)
    finally:
        await live.stop()
    await send_long(update, reply, rich=ctx.rich, markup=action_kb(ctx))
    await send_outbox(update, ctx)


async def cmd_bot_toggle(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """/botoff and /boton in a group: admins pause or resume the bot there."""
    chat = update.effective_chat
    if getattr(chat, "type", "private") == "private":
        await update.message.reply_text("Use /botoff or /boton inside a group (admins only).")
        return
    if not await tg_live.is_group_admin(context.bot, chat.id, update.effective_user.id):
        await update.message.reply_text("Only group admins can do that.")
        return
    off = update.message.text.lower().startswith("/botoff")
    tg_live.set_group_off(chat.id, off)
    await update.message.reply_text("Paused in this group. An admin can send /boton to bring me back."
                                    if off else "I'm back. Mention me or reply to me to ask something.")


# ---------- event viewership tracker (owner) ----------

EVENT_HELP = ("/event - live now, peak, average, hours watched, top channels\n"
              "/event chart - viewers-over-time chart\n"
              "/event page - full dashboard (open the file)\n"
              "/event top - all channels ranked\n"
              "/event add <YouTube link or @channel> - always track this stream or channel\n"
              "/event official <@channel or name> - mark a channel as official\n"
              "/event exclude <link, video id or channel name> - drop a stream that isn't about the event\n"
              "/event phrase <words> - also count streams whose title has these words\n"
              "/event start <name> | /event stop - start or pause tracking")


async def cmd_event(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    from core import events, eventpage
    args = context.args or []
    action = args[0].lower() if args else ""
    rest = " ".join(args[1:]).strip()
    ev = events.get()
    if action == "start":
        ev = events.create(rest.split()[0] if rest else "bmsd")
        await update.message.reply_text(f"Tracking {ev['name']}: searching YouTube for live streams every "
                                        f"{events.POLL_SEC}s.\nPhrases: {', '.join(ev['match'])}\n\n{EVENT_HELP}")
        return
    if not ev:
        await update.message.reply_text("No event yet. Start with: /event start bmsd")
        return
    key = ev["key"]
    if action == "stop":
        events.update(key, active=False)
        await update.message.reply_text(f"Paused {ev['name']}. Data is kept; /event start {key} resumes.")
    elif action in ("chart", "graph"):
        png = await asyncio.to_thread(eventpage.chart_png, key)
        if not png:
            await update.message.reply_text("No chart yet (no data, or matplotlib isn't installed: pip install matplotlib).")
            return
        import io as _io
        await update.message.reply_photo(photo=_io.BytesIO(png), caption=f"{ev['name']}: YouTube viewers over time")
    elif action in ("page", "dashboard", "html"):
        page = await asyncio.to_thread(eventpage.build_html, key)
        import io as _io
        await update.message.reply_document(document=_io.BytesIO(page.encode()), filename=f"{key}-viewership.html",
                                            caption="Open in your browser for the full dashboard (live data is a "
                                                    "snapshot; send /event page again for fresh numbers).")
    elif action == "top":
        st = await asyncio.to_thread(events.stats, key)
        live = {s["channel"] for s in st.get("live_now", [])}
        rows = [f"{i}. {c['channel']} ({c['kind']}{', LIVE' if c['channel'] in live else ', ended'}): "
                f"{events.fmt_n(c['hw'])} h watched, peak {events.fmt_n(c['peak'])}\n   {c['url']}"
                for i, c in enumerate(st.get("channels", [])[:60], 1)]
        await send_long(update, "\n".join(rows) or "No channels yet.", preview=False)
    elif action == "add" and rest:
        vid = re.search(r"(?:v=|youtu\.be/|/live/)([A-Za-z0-9_-]{11})", rest)
        if vid:
            events.update(key, videos=list(dict.fromkeys(ev.get("videos", []) + [vid.group(1)])))
            await update.message.reply_text(f"Tracking that stream ({vid.group(1)}) in {ev['name']}.")
        else:
            events.update(key, channels=list(dict.fromkeys(ev.get("channels", []) + [rest.split()[0]])))
            await update.message.reply_text(f"Tracking {rest.split()[0]}'s live streams in {ev['name']}.")
    elif action == "official" and rest:
        events.update(key, official=list(dict.fromkeys(ev.get("official", []) + [rest.lower().lstrip("@")])))
        await update.message.reply_text(f"'{rest}' counts as official from now on.")
    elif action in ("exclude", "remove") and rest:
        vid = re.search(r"(?:v=|youtu\.be/|/live/)([A-Za-z0-9_-]{11})", rest)
        item = vid.group(1) if vid else (rest if re.fullmatch(r"[A-Za-z0-9_-]{11}", rest) else rest.lower())
        events.update(key, excluded=list(dict.fromkeys(ev.get("excluded", []) + [item])))
        await update.message.reply_text(f"Excluded '{item}' from {ev['name']} from now on.")
    elif action == "phrase" and rest:
        events.update(key, match=list(dict.fromkeys(ev["match"] + [rest.lower()])),
                      queries=list(dict.fromkeys(ev["queries"] + [rest]))[:8])
        await update.message.reply_text(f"Also tracking streams with '{rest}' in the title.")
    elif action == "help":
        await update.message.reply_text(EVENT_HELP)
    else:
        text = await asyncio.to_thread(events.summary, key)
        await send_long(update, text, preview=False)


# ---------- Telegram channels ----------

async def cmd_channels(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """/channels, /channels add @a @b, /channels remove @a, /channels brief, /channels mute|unmute @a (owner)."""
    from core import channels
    user = await gate(update, count=False)
    if not user:
        return
    uid = user["uid"]
    args = context.args or []
    action = (args[0].lower() if args else "")
    names = args[1:]
    if action in ("add", "follow") and names:
        msgs = [await asyncio.to_thread(channels.follow, uid, n) for n in names[:5]]
        await update.message.reply_text("\n".join(msgs))
    elif action in ("remove", "unfollow", "rm") and names:
        await update.message.reply_text("\n".join(channels.unfollow(uid, n) for n in names))
    elif action in ("brief", "summary", "news"):
        if uid != "owner" and not gate_ok_for_ai(user):
            await update.message.reply_text(f"You've used your {DAILY_LIMIT} requests for today.")
            return
        await update.message.reply_text("Reading your channels...")
        try:
            text = await asyncio.to_thread(channels.channel_brief, uid)
            kb = None if text.startswith("Nothing new") else links_kb(uid)
        except Exception as e:
            text, kb = f"Couldn't build the channel brief: {str(e)[:200]}", None
        await send_long(update, text, markup=kb)
    elif action in ("mute", "unmute") and names and uid == "owner":
        muted = store.kv_get("tgchannels_muted", [])
        targets = [channels.clean_name(n).lower() or n.lower().lstrip("@") for n in names]
        muted = sorted(set(muted) | set(targets)) if action == "mute" else [m for m in muted if m not in targets]
        store.kv_set("tgchannels_muted", muted)
        await update.message.reply_text(("Muted: " if action == "mute" else "Unmuted: ") + ", ".join(targets))
    else:
        await update.message.reply_text(channels.status_text(uid))


def gate_ok_for_ai(user: dict) -> bool:
    if store.usage_today(user["uid"]) >= DAILY_LIMIT:
        return False
    store.use_one(user["uid"])
    return True


# ---------- inline mode: @bot <question> in any chat ----------

async def on_inline_query(update: Update, context: ContextTypes.DEFAULT_TYPE):
    iq = update.inline_query
    q = (iq.query or "").strip()
    uid = store.uid_for("telegram", str(iq.from_user.id))
    u = store.get_user(uid)
    if u and u.get("blocked"):
        await iq.answer([], cache_time=0, is_personal=True)
        return
    button = InlineQueryResultsButton(text="Open the full chat", start_parameter="inline")
    if len(q) < 3:
        await iq.answer([], cache_time=0, is_personal=True, button=button)
        return
    bot_name = context.bot.username or "bot"
    result = InlineQueryResultArticle(
        id=secrets.token_hex(8),
        title="Ask: " + (q if len(q) <= 60 else q[:60] + "…"),
        description="Tap to post the answer here",
        input_message_content=InputTextMessageContent(tg_live.inline_text(q, "<i>Thinking…</i>"), parse_mode="HTML"),
        reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton(f"Ask @{bot_name} yourself",
                                                                 url=f"https://t.me/{bot_name}?start=inline")]]),
    )
    await iq.answer([result], cache_time=0, is_personal=True, button=button)


async def on_inline_chosen(update: Update, context: ContextTypes.DEFAULT_TYPE):
    r = update.chosen_inline_result
    mid, q = r.inline_message_id, (r.query or "").strip()
    if not mid or not q:
        return
    bot_name = context.bot.username or "bot"
    kb = InlineKeyboardMarkup([[InlineKeyboardButton(f"Ask @{bot_name} yourself", url=f"https://t.me/{bot_name}?start=inline")]])

    async def show(body_html: str, plain: str = None):
        try:
            await context.bot.edit_message_text(inline_message_id=mid, text=tg_live.inline_text(q, body_html),
                                                parse_mode="HTML", disable_web_page_preview=True, reply_markup=kb)
        except Exception as e:
            if plain is not None and "not modified" not in str(e).lower():
                try:
                    await context.bot.edit_message_text(inline_message_id=mid, text=f"Q: {q}\n\n{plain}"[:4090],
                                                        reply_markup=kb)
                except Exception:
                    pass

    raw = str(r.from_user.id)
    uid = store.uid_for("telegram", raw)
    user = store.ensure_user(uid, "telegram", raw, r.from_user.first_name or "", r.from_user.username or "")
    if user.get("blocked"):
        await show("<i>Not available.</i>")
        return
    wait = tg_live.inline_cooldown(raw)
    if wait:
        await show(f"<i>Too fast. Try again in {wait}s.</i>")
        return
    if uid != "owner":
        if store.usage_today(uid) >= DAILY_LIMIT:
            await show(f"<i>Daily limit reached. Open @{html.escape(bot_name)} tomorrow for more.</i>")
            return
        store.use_one(uid)
    memory.clear(f"group:inline:{raw}", "telegram")  # every inline question stands alone
    ctx = agent_context.Ctx()
    task = asyncio.create_task(asyncio.to_thread(
        chat_fn, raw, "telegram", q, ctx, r.from_user.first_name or "", r.from_user.username or "", raw, "auto",
        None, {"id": f"inline:{raw}", "title": "inline", "inline": True}))
    shown = None
    while not task.done():
        await asyncio.sleep(1.5)
        label = (ctx.progress or "Thinking") + "…"
        if label != shown and not task.done():
            await show(f"<i>{html.escape(label)}</i>")
            shown = label
    try:
        answer = task.result()
    except Exception as e:
        answer = f"Something went wrong: {str(e)[:150]}"
    answer = answer if len(answer) < 3300 else answer[:3300].rsplit("\n", 1)[0] + "\n…"
    from core import tg_format
    extra = action_kb(ctx)
    if extra:
        kb = InlineKeyboardMarkup(list(extra.inline_keyboard) + list(kb.inline_keyboard))
    for item in ctx.outbox:  # files can't be posted inline; drop them
        try:
            os.unlink(item["path"])
        except OSError:
            pass
    await show(tg_format.to_html(answer), plain=answer)
    store.log_activity(uid, "telegram_inline", "inline", q)


async def on_my_membership(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Bot added to or removed from a group."""
    cm = update.my_chat_member
    chat = cm.chat
    if chat.type == "private":
        return
    groups = store.kv_get("groups", {})
    status = cm.new_chat_member.status
    who = cm.from_user
    by = f"@{who.username}" if who and who.username else (who.first_name if who else "someone")
    if status in ("member", "administrator"):
        groups[str(chat.id)] = {"title": chat.title, "added_by": by, "added_at": store.iso(store.now_utc())}
        store.kv_set("groups", groups)
        await notify_owner(context, f"I was added to the group '{chat.title}' ({chat.id}) by {by}. "
                                    f"Remove me with /leave {chat.id}")
        if GROUPS_ENABLED:
            try:
                await context.bot.send_message(chat_id=chat.id, text=GROUP_INTRO.format(bot=context.bot.username))
            except Exception:
                pass
    elif status in ("left", "kicked"):
        groups.pop(str(chat.id), None)
        store.kv_set("groups", groups)
        await notify_owner(context, f"I was removed from the group '{chat.title}' ({chat.id}).")


async def cmd_groups(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    groups = store.kv_get("groups", {})
    if not groups:
        await update.message.reply_text("I'm not in any groups. Add me to one from the group's member list.")
        return
    await send_long(update, "Groups I'm in:\n\n" + "\n".join(
        f"{g['title']} ({gid}), added by {g.get('added_by')}" for gid, g in groups.items())
        + "\n\nLeave one with /leave <id>")


async def cmd_leave(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update) or not context.args:
        return
    gid = context.args[0]
    try:
        await context.bot.leave_chat(int(gid))
        groups = store.kv_get("groups", {})
        groups.pop(gid, None)
        store.kv_set("groups", groups)
        await update.message.reply_text(f"Left group {gid}.")
    except Exception as e:
        await update.message.reply_text(f"Couldn't leave: {str(e)[:150]}")


# ---------- content studio and backup (owner) ----------

async def _owner_prompt(update, context, prompt: str):
    user = await gate(update)
    if user:
        await execute_chat(update, context, prompt, session_model(user) or "sonnet")


async def cmd_script(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    topic = " ".join(context.args).strip()
    if not topic:
        await update.message.reply_text("Usage: /script <topic or idea>. Add your style first with /brandvoice add <sample>.")
        return
    await _owner_prompt(update, context, (
        f"Write a video script about: {topic}\n\nFirst call get_brand_voice and match it. Give: 3 hook options for the "
        "first 3 seconds; the full script with retention beats (pattern interrupts every 20-30 seconds) and on-screen "
        "text cues; a call to action; 5 title options; 3 thumbnail concepts (text overlay + visual). Say which format "
        "fits best (Short, Reel or long-form) and why. Plain text, no markdown."))


async def cmd_shorts(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    url = " ".join(context.args).strip()
    if not url:
        await update.message.reply_text("Usage: /shorts <YouTube link>. I'll find 5 clip moments with timestamps.")
        return
    await _owner_prompt(update, context, (
        f"Turn this long video into Shorts/Reels: {url}\n\nCall video_transcript, then get_brand_voice. Give 5 clips, "
        "each with: start and end timestamps (30-60 seconds), the hook as on-screen text, why it will hold attention, a "
        "caption, and 5 hashtags. Rank them by viral potential. Plain text, no markdown."))


async def cmd_captions(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    topic = " ".join(context.args).strip()
    if not topic:
        await update.message.reply_text("Usage: /captions <what the post is about>")
        return
    await _owner_prompt(update, context, (
        f"Write a caption pack for a post about: {topic}\n\nCall get_brand_voice first. Give one caption each for "
        "Instagram (with 10 relevant hashtags, mixing big and niche), YouTube Shorts (title + description), X (under "
        "280 characters) and LinkedIn (professional angle). Plain text, no markdown."))


async def cmd_brandvoice(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    from core import scout
    args = " ".join(context.args).strip()
    if args.lower().startswith("add "):
        await update.message.reply_text(scout.save_brand_voice(args[4:]))
    elif args.lower() == "clear":
        store.kv_set("brand_voice", [])
        await update.message.reply_text("Brand voice samples cleared.")
    else:
        n = len(store.kv_get("brand_voice", []))
        await update.message.reply_text(
            f"I have {n} brand voice sample(s).\n\nAdd one: /brandvoice add <paste a script, caption or post you wrote>\n"
            "The more you add (up to 10), the more scripts and captions sound like you. Clear: /brandvoice clear")


async def cmd_backup(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    from core import backup
    try:
        path = await asyncio.to_thread(backup.run_backup)
        await update.message.reply_text(f"Backed up to {path}\n(Last {backup.KEEP} copies are kept; nightly at 3 AM.)")
    except Exception as e:
        await update.message.reply_text(f"Backup failed: {str(e)[:200]}")


async def cmd_scout(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update):
        return
    handle = " ".join(context.args).strip()
    if not handle:
        await update.message.reply_text("Usage: /scout @channel or a YouTube/Instagram link")
        return
    await _owner_prompt(update, context, f"Scout this creator with scout_creator and give me a fit assessment for my "
                                         f"talent agency (strengths, risks, rate range to expect, brand categories): {handle}")


# What people see on the bot's profile ("About") and on the empty chat screen before they press Start.
BOT_ABOUT = os.getenv("BOT_ABOUT") or "Your AI assistant on Telegram. Ask anything, anytime."
BOT_DESCRIPTION = os.getenv("BOT_DESCRIPTION") or (
    "Your AI assistant on Telegram.\n\n"
    "Ask questions, get answers with sources, summarise links and files, and stay updated on what you care about.\n\n"
    "Free to use. Tap Start.")
PUBLIC_COMMANDS = [("new", "Start a fresh conversation"), ("search", "Web answer with sources"),
                   ("news", "Your news now"), ("interests", "Pick topics for updates"),
                   ("updates", "How often you get updates"), ("channels", "Follow Telegram channels"),
                   ("alert", "Alert me when a keyword shows up"), ("model", "Choose the AI model"),
                   ("style", "How you like replies"), ("profile", "What I know about you"),
                   ("forget", "Erase what I know about you"), ("help", "Everything I can do")]
OWNER_COMMANDS = [("event", "Tournament viewership"), ("brief", "Full briefing now"), ("radar", "Content and business radar"),
                  ("status", "Health check"), ("usage", "Costs today and this week"), ("users", "Who's using the bot")]


async def setup_profile(app):
    """Set the About text, description and command menu (only when they changed)."""
    from telegram import BotCommand, BotCommandScopeAllPrivateChats, BotCommandScopeChat
    bot = app.bot
    try:
        about, desc = BOT_ABOUT[:120], BOT_DESCRIPTION.replace("{bot}", bot.username or "this bot")[:512]
        if (await bot.get_my_short_description()).short_description != about:
            await bot.set_my_short_description(about)
        if (await bot.get_my_description()).description != desc:
            await bot.set_my_description(desc)
        await bot.set_my_commands([BotCommand(c, d) for c, d in PUBLIC_COMMANDS], scope=BotCommandScopeAllPrivateChats())
        for oid in store.owner_ids()["telegram"]:
            await bot.set_my_commands([BotCommand(c, d) for c, d in PUBLIC_COMMANDS + OWNER_COMMANDS],
                                      scope=BotCommandScopeChat(int(oid)))
    except Exception as e:
        print(f"Couldn't update the bot profile: {str(e)[:150]}")


def run_telegram():
    token = os.getenv("TELEGRAM_BOT_TOKEN")
    app = ApplicationBuilder().token(token).concurrent_updates(True).post_init(setup_profile).build()
    PRIVATE = filters.ChatType.PRIVATE
    GROUPS = filters.ChatType.GROUPS
    app.add_handler(CommandHandler("start", cmd_start, filters=PRIVATE))
    app.add_handler(CommandHandler("help", cmd_help, filters=PRIVATE))
    app.add_handler(CommandHandler(["clear", "new"], cmd_clear, filters=PRIVATE))
    app.add_handler(CommandHandler("search", cmd_search, filters=PRIVATE))
    app.add_handler(CommandHandler("profile", cmd_profile, filters=PRIVATE))
    app.add_handler(CommandHandler("forget", cmd_forget, filters=PRIVATE))
    app.add_handler(CommandHandler("interests", cmd_interests, filters=PRIVATE))
    app.add_handler(CommandHandler(["newson", "newsoff"], cmd_news_toggle, filters=PRIVATE))
    app.add_handler(CommandHandler("keywords", cmd_keywords, filters=PRIVATE))
    app.add_handler(CommandHandler("news", cmd_news, filters=PRIVATE))
    app.add_handler(CommandHandler("brief", cmd_brief, filters=PRIVATE))
    app.add_handler(CommandHandler("radar", cmd_radar, filters=PRIVATE))
    app.add_handler(CommandHandler("users", cmd_users, filters=PRIVATE))
    app.add_handler(CommandHandler("status", cmd_status, filters=PRIVATE))
    app.add_handler(CommandHandler("usage", cmd_usage, filters=PRIVATE))
    app.add_handler(CommandHandler("credits", cmd_credits, filters=PRIVATE))
    app.add_handler(CommandHandler("dashboard", cmd_dashboard, filters=PRIVATE))
    app.add_handler(CommandHandler("pipeline", cmd_pipeline, filters=PRIVATE))
    app.add_handler(CommandHandler("alert", cmd_alert, filters=PRIVATE))
    app.add_handler(CommandHandler("alerts", cmd_alerts, filters=PRIVATE))
    app.add_handler(CommandHandler("unalert", cmd_unalert, filters=PRIVATE))
    app.add_handler(CommandHandler(["block", "unblock"], cmd_block, filters=PRIVATE))
    app.add_handler(CommandHandler(["deep", "quick"], cmd_model, filters=PRIVATE))
    app.add_handler(CommandHandler("model", cmd_model_default, filters=PRIVATE))
    app.add_handler(CommandHandler("picker", cmd_picker, filters=PRIVATE))
    app.add_handler(CommandHandler("switch", cmd_switch, filters=PRIVATE))
    app.add_handler(CommandHandler("style", cmd_style, filters=PRIVATE))
    app.add_handler(CommandHandler("thinking", cmd_thinking, filters=PRIVATE))
    app.add_handler(CommandHandler("think", cmd_think, filters=PRIVATE))
    app.add_handler(CommandHandler("activity", cmd_activity, filters=PRIVATE))
    app.add_handler(CallbackQueryHandler(on_model_pick, pattern=r"^m:"))
    app.add_handler(CallbackQueryHandler(on_topic_pick, pattern=r"^(t|tx|f|z):"))
    app.add_handler(CallbackQueryHandler(on_download_pick, pattern=r"^dl:"))
    app.add_handler(CallbackQueryHandler(on_news_links, pattern=r"^nl:"))
    app.add_handler(CallbackQueryHandler(on_file_action, pattern=r"^fa:"))
    app.add_handler(CommandHandler("channels", cmd_channels, filters=PRIVATE))
    app.add_handler(CommandHandler("event", cmd_event, filters=PRIVATE))
    app.add_handler(InlineQueryHandler(on_inline_query))
    app.add_handler(ChosenInlineResultHandler(on_inline_chosen))
    app.add_handler(CommandHandler(["botoff", "boton"], cmd_bot_toggle))
    app.add_handler(CommandHandler("updates", cmd_updates, filters=PRIVATE))
    app.add_handler(MessageHandler(PRIVATE & filters.TEXT & ~filters.COMMAND, handle_message))
    app.add_handler(MessageHandler(PRIVATE & (filters.VOICE | filters.AUDIO), handle_voice))
    app.add_handler(MessageHandler(PRIVATE & filters.PHOTO, handle_photo))
    app.add_handler(MessageHandler(PRIVATE & filters.Document.ALL, handle_document))
    app.add_handler(CommandHandler("help", cmd_help_group, filters=GROUPS))
    app.add_handler(CommandHandler("ask", cmd_ask))
    app.add_handler(CommandHandler("groups", cmd_groups, filters=PRIVATE))
    app.add_handler(CommandHandler("leave", cmd_leave, filters=PRIVATE))
    app.add_handler(CommandHandler("script", cmd_script, filters=PRIVATE))
    app.add_handler(CommandHandler("shorts", cmd_shorts, filters=PRIVATE))
    app.add_handler(CommandHandler("captions", cmd_captions, filters=PRIVATE))
    app.add_handler(CommandHandler("brandvoice", cmd_brandvoice, filters=PRIVATE))
    app.add_handler(CommandHandler("backup", cmd_backup, filters=PRIVATE))
    app.add_handler(CommandHandler("scout", cmd_scout, filters=PRIVATE))
    app.add_handler(MessageHandler(GROUPS & (filters.TEXT | filters.CAPTION) & ~filters.COMMAND, handle_group))
    app.add_handler(ChatMemberHandler(on_my_membership, ChatMemberHandler.MY_CHAT_MEMBER))
    app.add_error_handler(on_error)
    print(f"Telegram: running (open to everyone, {DAILY_LIMIT} requests/day for non-owners)")
    app.run_polling(allowed_updates=Update.ALL_TYPES)
