import os
import asyncio
import discord
import tempfile
import fitz
import base64
import requests
from pptx import Presentation
from docx import Document
import openpyxl
from core.agent import chat
from core.memory import Memory
from core import store
from dotenv import load_dotenv
import anthropic

load_dotenv()

memory = Memory()
client = anthropic.Anthropic()
ALLOWED = set(os.getenv("ALLOWED_DISCORD_IDS", "").split(","))

intents = discord.Intents.default()
intents.message_content = True
bot = discord.Client(intents=intents)


def extract_pdf(path: str) -> str:
    doc = fitz.open(path)
    text = ""
    for page in doc:
        text += page.get_text()
    return text[:8000]


def extract_pptx(path: str) -> str:
    prs = Presentation(path)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text[:8000]


def extract_docx(path: str) -> str:
    doc = Document(path)
    text = "\n".join([p.text for p in doc.paragraphs])
    return text[:8000]


def extract_txt(path: str) -> str:
    with open(path, "r", errors="ignore") as f:
        return f.read()[:8000]


def extract_excel(path: str) -> str:
    wb = openpyxl.load_workbook(path)
    text = ""
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        text += f"Sheet: {sheet}\n"
        for row in ws.iter_rows(values_only=True):
            row_text = " | ".join([str(cell) for cell in row if cell is not None])
            if row_text.strip():
                text += row_text + "\n"
    return text[:8000]


def summarize_text(user_id: str, text: str, filename: str) -> str:
    message = f"I am sending you the content of a file called '{filename}'. Please read it and summarize the key points, insights, and anything important.\n\nFile content:\n{text}"
    return chat(user_id, "discord", message)


def summarize_image(path: str) -> str:
    with open(path, "rb") as f:
        image_data = f.read()
    ext = path.split(".")[-1].lower()
    media_map = {
        "jpg": "image/jpeg", "jpeg": "image/jpeg",
        "png": "image/png", "gif": "image/gif", "webp": "image/webp"
    }
    media_type = media_map.get(ext, "image/jpeg")
    b64 = base64.standard_b64encode(image_data).decode("utf-8")
    response = client.messages.create(
        model="claude-sonnet-5-5",
        max_tokens=1024,
        messages=[{
            "role": "user",
            "content": [
                {
                    "type": "image",
                    "source": {"type": "base64", "media_type": media_type, "data": b64}
                },
                {
                    "type": "text",
                    "text": "Describe this image in detail. If there is any text, extract it. Summarize everything important."
                }
            ]
        }]
    )
    return response.content[0].text


def generate_image(prompt: str) -> bytes:
    encoded = prompt.replace(" ", "%20")
    url = f"https://image.pollinations.ai/prompt/{encoded}?width=1024&height=1024&nologo=true"
    response = requests.get(url, timeout=60)
    return response.content


def is_image_request(text: str) -> bool:
    triggers = [
        "generate image", "create image", "make image", "draw",
        "generate a picture", "create a picture", "make a picture",
        "imagine", "visualize", "generate photo", "create photo"
    ]
    return any(t in text.lower() for t in triggers)


def extract_image_prompt(text: str) -> str:
    triggers = [
        "generate image of", "create image of", "make image of",
        "generate a picture of", "create a picture of", "make a picture of",
        "generate photo of", "create photo of", "draw", "imagine", "visualize",
        "generate image", "create image", "make image"
    ]
    text_lower = text.lower()
    for t in triggers:
        if t in text_lower:
            return text_lower.split(t)[-1].strip()
    return text


@bot.event
async def on_ready():
    print(f"Discord: logged in as {bot.user}")


@bot.event
async def on_message(message):
    if message.author == bot.user:
        return

    if str(message.author.id) not in ALLOWED:
        return

    is_dm = isinstance(message.channel, discord.DMChannel)
    is_mentioned = bot.user in message.mentions

    if not (is_dm or is_mentioned):
        return

    user_id = str(message.author.id)
    content = message.content.replace(f"<@{bot.user.id}>", "").strip()

    if content.lower() == "!clear":
        memory.clear(store.uid_for("discord", user_id), "discord")
        await message.channel.send("Memory cleared.")
        return

    if content.lower() in ["!brief", "brief me", "send brief"]:
        await message.channel.send("Building your briefing now. Give me 2-3 minutes...")
        from digest import build_digest
        digest = await asyncio.to_thread(build_digest)
        chunks = [digest[i:i+2000] for i in range(0, len(digest), 2000)]
        for chunk in chunks:
            await message.channel.send(chunk)
        return

    if message.attachments:
        for attachment in message.attachments:
            filename = attachment.filename.lower()
            await message.channel.send("Reading your file...")

            file_bytes = await attachment.read()
            ext = os.path.splitext(filename)[1]

            with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
                tmp.write(file_bytes)
                tmp_path = tmp.name

            try:
                if filename.endswith(".pdf"):
                    text = extract_pdf(tmp_path)
                    response = await asyncio.to_thread(summarize_text, user_id, text, attachment.filename)
                elif filename.endswith((".pptx", ".ppt")):
                    text = extract_pptx(tmp_path)
                    response = await asyncio.to_thread(summarize_text, user_id, text, attachment.filename)
                elif filename.endswith((".docx", ".doc")):
                    text = extract_docx(tmp_path)
                    response = await asyncio.to_thread(summarize_text, user_id, text, attachment.filename)
                elif filename.endswith((".txt", ".csv")):
                    text = extract_txt(tmp_path)
                    response = await asyncio.to_thread(summarize_text, user_id, text, attachment.filename)
                elif filename.endswith((".xlsx", ".xls")):
                    text = extract_excel(tmp_path)
                    response = await asyncio.to_thread(summarize_text, user_id, text, attachment.filename)
                elif filename.endswith((".jpg", ".jpeg", ".png", ".gif", ".webp")):
                    response = await asyncio.to_thread(summarize_image, tmp_path)
                else:
                    response = "File type not supported. Send PDF, Word, PowerPoint, Excel, image, or text files."
            except Exception as e:
                response = f"Could not read file: {str(e)}"
            finally:
                os.unlink(tmp_path)

            if len(response) <= 2000:
                await message.channel.send(response)
            else:
                for i in range(0, len(response), 1990):
                    await message.channel.send(response[i:i+1990])
        return

    if not content:
        return

    if is_image_request(content):
        prompt = extract_image_prompt(content)
        await message.channel.send("Generating image...")
        try:
            image_bytes = await asyncio.to_thread(generate_image, prompt)
            with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as tmp:
                tmp.write(image_bytes)
                tmp_path = tmp.name
            await message.channel.send(file=discord.File(tmp_path, filename="generated.png"))
            os.unlink(tmp_path)
        except Exception as e:
            await message.channel.send(f"Could not generate image: {str(e)}")
        return

    async with message.channel.typing():
        response = await asyncio.to_thread(chat, user_id, "discord", content)

    if len(response) <= 2000:
        await message.channel.send(response)
    else:
        for i in range(0, len(response), 1990):
            await message.channel.send(response[i:i+1990])


def run_discord():
    bot.run(os.getenv("DISCORD_BOT_TOKEN"))